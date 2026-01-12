import os
from pathlib import Path
from typing import Dict, List, Any
import yaml
from pptx import Presentation


def _extract_text_from_shape(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return "\n".join(
        [run.text for para in shape.text_frame.paragraphs for run in para.runs]
    )


def _is_highlighted(shape) -> bool:
    if hasattr(shape, "fill") and shape.fill.type != 0:
        return True
    return False


def _extract_media(
    shape, slide_idx: int, shape_idx: int, media_dir: Path
) -> str | None:
    if not hasattr(shape, "image"):
        return None

    try:
        image_bytes = shape.image.blob
        ext = shape.image.ext
        filename = f"slide_{slide_idx}_shape_{shape_idx}{ext}"
        filepath = media_dir / filename

        with open(filepath, "wb") as f:
            f.write(image_bytes)

        return f"media/{filename}"
    except Exception:
        return None


def ppt_to_yaml(pptx_path: str, yaml_output_dir: str) -> str:
    pptx_file = Path(pptx_path)
    yaml_dir = Path(yaml_output_dir)
    yaml_dir.mkdir(parents=True, exist_ok=True)

    media_dir = yaml_dir / "media"
    media_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(pptx_file))

    slides_data = []
    highlighted_sections = []

    for slide_idx, slide in enumerate(prs.slides):
        slide_data = {
            "slide_number": slide_idx + 1,
            "title": "",
            "content": [],
            "media": [],
            "is_highlighted": False,
            "layout": "",
        }

        layout_name = slide.slide_layout.name if slide.slide_layout else ""
        slide_data["layout"] = layout_name

        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                text = _extract_text_from_shape(shape)
                if text.strip():
                    is_title = shape_idx == 0 and text.strip()
                    if is_title and not slide_data["title"]:
                        slide_data["title"] = text.strip()
                    else:
                        slide_data["content"].append(
                            {"type": "text", "value": text.strip()}
                        )

            if hasattr(shape, "image"):
                media_path = _extract_media(shape, slide_idx, shape_idx, media_dir)
                if media_path:
                    slide_data["media"].append({"type": "image", "path": media_path})

            if _is_highlighted(shape):
                slide_data["is_highlighted"] = True

        if slide_data["title"] or slide_data["content"] or slide_data["media"]:
            slides_data.append(slide_data)
            if slide_data["is_highlighted"]:
                highlighted_sections.append(
                    {
                        "slide_number": slide_idx + 1,
                        "title": slide_data["title"],
                        "content": slide_data["content"][:3],
                    }
                )

    output_data = {
        "title": pptx_file.stem,
        "hero_image": None,
        "slides": slides_data,
        "highlighted_sections": highlighted_sections,
        "total_slides": len(slides_data),
    }

    yaml_path = yaml_dir / f"{pptx_file.stem}.yaml"
    with open(yaml_path, "w", encoding="utf-8") as f:
        yaml.dump(
            output_data,
            f,
            allow_unicode=True,
            default_flow_style=False,
            sort_keys=False,
        )

    return str(yaml_path)
